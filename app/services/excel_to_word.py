#!/usr/bin/env python3
"""
excel_to_word.py — цифры из Excel-отчёта (форма № 230) -> текст доклада -> оформленный Word.
"""
import argparse
import os
import re
import sys
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# =============================================================================
# 1. НАСТРОЙКИ: какие показатели искать и как их называть в тексте
# =============================================================================
COLUMN_HINT = "в сфере восстановления"   # фрагмент названия графы
OPTIONAL = set()   # Строки, которые есть не во всех формах

# ключ, фрагменты названия строки (все должны входить в подпись), единица, название для таблицы
INDICATORS = [
    ("messages", ["Поступило сообщений о преступлениях"], "ед.", "Поступило сообщений"),
    ("refused", ["отказано в возбуждении дел"], "ед.", "Отказано в возбуждении дел"),
    ("cancelled", ["Отменено постановлений о возбуждении уголовного дела"], "ед.", "Отменено постановлений о возбуждении УД"),
    ("cancelled_refused", ["Отменено постановлений об отказе в возбуждении"], "ед.", "Отменено постановлений об отказе в возбуждении"),
    ("in_progress", ["Находилось в производстве уголовных дел"], "ед.", "Находилось в производстве УД"),
    ("to_court", ["Направлено уголовных дел в суд"], "ед.", "Направлено уголовных дел в суд"),
    ("terminated", ["Прекращено уголовных дел"], "ед.", "Прекращено уголовных дел"),
    ("houses", ["Число домов, признанных аварийными"], "ед.", "Число аварийных домов"),
    ("citizens", ["Число граждан, не переселенных из аварийного"], "чел.", "Число граждан, не переселённых из аварийного жилья"),
    ("court", ["Количество неисполненных судебных решений"], "ед.", "Неисполненные судебные решения"),
    ("funds", ["Размер денежных средств, выделенных на расселение"], "руб.", "Выделено средств"),
    ("unspent", ["Размер неосвоенных денежных средств, выделенных на переселение"], "тыс. руб.", "Не освоено средств"),
]

TITLE = (
    "Рассмотрение сообщений и расследование уголовных дел о преступлениях в сфере "
    "восстановления прав граждан, проживающих в аварийном жилье"
)

DEFAULT_TEXT = [
    "За {{period}} в следственное управление поступило {{messages}} сообщений о преступлениях "
    "в сфере восстановления прав граждан, проживающих в аварийном жилье[[ (АППГ – {{messages_appg_num}})]].",

    "Решения об отказе в возбуждении уголовного дела принимались в {{refused}} случаях.",

    "Отменено {{cancelled}} постановлений о возбуждении уголовного дела[[ (АППГ – {{cancelled_appg_num}})]]. "
    "Постановления об отказе в возбуждении уголовного дела отменялись {{cancelled_refused}} раз.",

    "В производстве следователей следственного управления находилось {{in_progress}} "
    "уголовных дел о преступлениях данной категории[[ (АППГ – {{in_progress_appg_num}})]].",

    "Направлено в суд {{to_court}} уголовных дел[[ (АППГ – {{to_court_appg_num}})]].",

    "Прекращено {{terminated}} уголовных дел о преступлениях данной категории[[ (АППГ – {{terminated_appg_num}})]].",

    "На территории Челябинской области числится {{houses}} домов, признанных аварийными, "
    "и {{citizens}} граждан, не переселённых из аварийного жилищного фонда. "
    "Количество неисполненных судебных решений — {{court}}. "
    "На расселение выделено {{funds}}.",
]

# =============================================================================
# 2. ЧТЕНИЕ EXCEL
# =============================================================================
CONNECTORS = ("в том числе", "в т.ч.", "из них", "по ним")
MONTH_NOM = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август",
             "сентябрь", "октябрь", "ноябрь", "декабрь"]
MONTH_GEN = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
             "сентября", "октября", "ноября", "декабря"]


def clean(s):
    s = str(s).replace("\xa0", " ")
    s = re.sub(r"(\w)-\s+(\w)", r"\1\2", s)
    return re.sub(r"\s+", " ", s).strip()


def key(s):
    s = clean(s).lower().replace("ё", "е")
    s = re.sub(r"\((?:из|по данным)?\s*стр[^)]*\)", "", s)
    return re.sub(r"[^a-zа-я0-9]", "", s)


def is_connector(s):
    t = re.sub(r"\(.*?\)", "", clean(s).lower()).strip(" :.")
    return any(t.startswith(c) and len(t) <= len(c) + 12 for c in CONNECTORS)


def to_num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if s in ("", "-", "–", "—", "х", "x", "Х", "X"):
        return None
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def read_sheets(path):
    if path.lower().endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(path, logfile=open(os.devnull, "w"))
        return [[[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
                for sh in wb.sheets()]
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return [[["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets]


class Report:
    def __init__(self, path, column_hint=COLUMN_HINT):
        self.path = path
        self.warnings = []
        for grid in read_sheets(path):
            w = max((len(r) for r in grid), default=0)
            self.g = [list(r) + [""] * (w - len(r)) for r in grid]
            if self._find_layout():
                break
        else:
            raise ValueError(f"{path}: не найдена таблица формы (строка с номерами граф 1, 2, 3…)")
        self._read_rows()
        self._pick_column(column_hint)
        self.period = self._period()

    def cell(self, r, c):
        return self.g[r][c] if 0 <= r < len(self.g) and 0 <= c < len(self.g[r]) else ""

    def _find_layout(self):
        for r in range(min(len(self.g), 60)):
            for c0 in range(1, len(self.g[r])):
                if to_num(self.cell(r, c0)) != 1:
                    continue
                n = 1
                while to_num(self.cell(r, c0 + n)) == n + 1:
                    n += 1
                if any(to_num(self.cell(rr, c0 - 1)) == 1 for rr in range(r + 1, r + 5)):
                    self.numbers_row, self.rownum_col = r, c0 - 1
                    self.data_cols = list(range(c0, c0 + n))
                    return True
        return False

    def _read_rows(self):
        L = self.rownum_col
        levels, parent, self.rows = [""] * L, "", []
        for r in range(self.numbers_row + 1, len(self.g)):
            row_no = to_num(self.cell(r, self.rownum_col))
            if row_no is None:
                continue
            cells = [clean(self.cell(r, j)) for j in range(L)]
            first = next((j for j, t in enumerate(cells) if t), None)
            if first is not None:
                levels[first:] = cells[first:]
                if first == 0 and not is_connector(cells[0]):
                    parent = cells[0]
            label = " / ".join(p for p in [parent] + [t for t in levels[1:] if t and not is_connector(t)] if p)
            self.rows.append({"r": r, "row_no": row_no, "label": label, "key": key(label)})

    def _pick_column(self, hint):
        for c in self.data_cols:
            for r in range(max(0, self.numbers_row - 7), self.numbers_row):
                if hint and hint in clean(self.cell(r, c)).lower():
                    self.col, self.col_title = c, clean(self.cell(r, c))
                    return
        self.col, self.col_title = self.data_cols[-1], f"графа {len(self.data_cols)}"
        self.warnings.append(f"Графа «{hint}» не найдена по названию — взята последняя ({self.col_title}).")

    def _period(self):
        m = re.search(r"_(\d{2})(\d{2})\.xlsx?$", os.path.basename(self.path).strip(), re.I)
        return (2000 + int(m.group(1)), int(m.group(2))) if m else None

    def find(self, fragments):
        ks = [key(f) for f in fragments]
        hits = [x for x in self.rows if all(k in x["key"] for k in ks)]
        return min(hits, key=lambda x: (x["label"].count(" / "), len(x["key"]))) if hits else None

    def values(self):
        out = {}
        for k, frags, unit, title in INDICATORS:
            row = self.find(frags)
            if row is None:
                out[k] = {"value": None, "row_no": None, "label": None}
                if k not in OPTIONAL:
                    self.warnings.append(f"Не найдена строка «{frags[-1]}» — будет «нет данных».")
            else:
                out[k] = {"value": to_num(self.cell(row["r"], self.col)),
                          "row_no": row["row_no"], "label": row["label"]}
        return out


def find_report(search_dirs, code, form, year, month):
    pat = re.compile(rf"^\s*{code}_{form}_{year % 100:02d}{month:02d}\s*\.xlsx?$", re.I)
    for d in search_dirs:
        for dp, _, files in os.walk(d):
            for f in sorted(files):
                if pat.match(f):
                    return os.path.join(dp, f)
    return None

# =============================================================================
# 3. ЧИСЛА -> ТЕКСТ
# =============================================================================
NBSP = " "

def fnum(v):
    if v is None:
        return "<нет данных>"
    if v < 0:
        return "−" + fnum(-v)
    s = f"{v:,.1f}" if isinstance(v, float) and not v.is_integer() else f"{int(round(v)):,}"
    return s.replace(",", NBSP).replace(".", ",")


def fpct(p, signed=False):
    if p is None:
        return "—"
    sign = ("+" if p > 0 else "−" if p < 0 else "") if signed else ""
    return f"{sign}{abs(p):.1f}".replace(".", ",") + NBSP + "%"


def change(cur, base):
    if cur is None or base is None:
        return None, None
    d = cur - base
    return d, (round(d / base * 100, 1) if base else None)


def period_label(y, m):
    return f"январь {y} г." if m == 1 else f"январь–{MONTH_NOM[m - 1]} {y} г."


def build_context(cur, cur_period, bases):
    y, m = cur_period
    ctx = {"period": period_label(y, m), "period_end": f"конец {MONTH_GEN[m - 1]} {y} г.", "year": str(y)}
    for k, _, unit, _ in INDICATORS:
        v = cur[k]["value"]
        ctx[k] = None if v is None else f"{fnum(v)} {unit}"
        ctx[f"{k}_num"] = None if v is None else fnum(v)
        for tag in ("prev", "appg"):
            vals, per = bases.get(tag, (None, None))
            b = vals[k]["value"] if vals else None
            d, p = change(v, b)
            ctx[f"{k}_{tag}_num"] = None if b is None else fnum(b)
            ctx[f"{k}_{tag}_delta"] = None if d is None else ("+" if d > 0 else "") + fnum(d)
            ctx[f"{k}_{tag}_pct"] = None if p is None else fpct(p, signed=True)
            if d is None:
                ctx[f"{k}_{tag}"] = None
                continue
            base_txt = f"{fnum(b)} {unit}"
            trend = ("без изменений" if d == 0 else
                     f"{'рост' if d > 0 else 'снижение'} на {fnum(abs(d))} {unit}"
                     + (f", или на {fpct(p)}" if p is not None else ""))
            if tag == "appg":
                ctx[f"{k}_{tag}"] = f"на уровне АППГ – {base_txt}" if d == 0 else f"АППГ – {base_txt}; {trend}"
            else:
                ctx[f"{k}_{tag}"] = f"{trend} (на конец {MONTH_GEN[per[1] - 1]} {per[0]} г. – {base_txt})"
    return ctx


PLACEHOLDER = re.compile(r"\{\{\s*(\w+)\s*\}\}")
UNKNOWN = set()


def render(text, ctx):
    out = []
    for part in re.split(r"(\[\[.*?\]\])", text, flags=re.S):
        optional = part.startswith("[[") and part.endswith("]]")
        body = part[2:-2] if optional else part
        segs, complete, pos = [], True, 0
        for m in PLACEHOLDER.finditer(body):
            segs.append((body[pos:m.start()], False))
            name = m.group(1)
            v = ctx.get(name)
            if v is None:
                complete = False
                if name not in ctx:
                    UNKNOWN.add(name)
                v = "нет данных" if name in ctx else "{{" + name + "}}"
            segs.append((v, True))
            pos = m.end()
        segs.append((body[pos:], False))
        if optional and not complete:
            continue
        out += segs
    merged = [[t, h] for t, h in out if t]
    for a, b in zip(merged, merged[1:]):
        if a[0].endswith(".") and b[0].startswith("."):
            b[0] = b[0][1:]
    return [(t, h) for t, h in merged if t]

# =============================================================================
# 4. WORD
# =============================================================================
FONT = "Times New Roman"

def setup_document(doc):
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
        add_page_number(s)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(14)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(11)


def add_runs(paragraph, segs, highlight=True, like=None):
    for text, is_value in segs:
        r = paragraph.add_run(text)
        if like is not None:
            r.font.name, r.font.size = like.font.name, like.font.size
            r.font.bold, r.font.italic, r.font.underline = like.font.bold, like.font.italic, like.font.underline
        if is_value and highlight:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def body_paragraph(doc, segs, highlight):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    add_runs(p, segs, highlight)
    return p


def centered(doc, text, size=14, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def summary_table(doc, cur, bases, cur_period):
    prev_vals = bases.get("prev", (None, None))[0]
    appg_vals = bases.get("appg", (None, None))[0]
    prev_per = bases.get("prev", (None, None))[1]
    appg_per = bases.get("appg", (None, None))[1]
    prev_head = f"Пред. период ({MONTH_NOM[prev_per[1] - 1]} {prev_per[0]})" if prev_per else "Пред. период"
    appg_head = f"АППГ ({MONTH_NOM[appg_per[1] - 1]} {appg_per[0]})" if appg_per else "АППГ"
    header = ["Показатель", "Ед.", period_label(*cur_period), prev_head, "Δ", "%", appg_head, "Δ", "%"]
    widths = [3.8, 1.0, 1.9, 1.8, 1.6, 1.5, 1.8, 1.6, 1.4]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trPr.append(hdr)
    for k, _, unit, title in INDICATORS:
        if cur[k]["row_no"] is None and k in OPTIONAL:
            continue
        v = cur[k]["value"]
        row = [title, unit, fnum(v)]
        for vals in (prev_vals, appg_vals):
            b = vals[k]["value"] if vals else None
            d, p = change(v, b)
            row += [fnum(b) if vals else "нет отчёта", "—" if d is None else ("+" if d > 0 else "") + fnum(d),
                    fpct(p, signed=True)]
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = txt
    for i, txt in enumerate(header):
        t.rows[0].cells[i].text = txt
        shade(t.rows[0].cells[i], "D9E2F3")
    t.autofit = False
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "57")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    t._tbl.tblPr.append(mar)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 567)))
    for ri, r in enumerate(t.rows):
        cant = OxmlElement("w:cantSplit")
        r._tr.get_or_add_trPr().append(cant)
        for i, c in enumerate(r.cells):
            c.width = Cm(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Cm(0)
                if i > 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.bold = ri == 0
    return t


def build_default(cur, bases, cur_period, ctx, highlight, sources):
    doc = Document()
    setup_document(doc)
    centered(doc, TITLE, bold=True)
    centered(doc, f"(по данным статистической формы № 230 за {ctx['period']})", size=12, italic=True)
    for text in DEFAULT_TEXT:
        segs = render(text, ctx)
        if segs:
            body_paragraph(doc, segs, highlight)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Таблица. Показатели раздела и их динамика")
    r.font.size = Pt(12)
    summary_table(doc, cur, bases, cur_period)
    footer_note(doc, sources)
    return doc


def footer_note(doc, sources):
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(f"Сформировано автоматически {date.today():%d.%m.%Y} из файлов: " + "; ".join(sources) +
                  ". Подставленные значения выделены жёлтым; пустые ячейки отчёта указаны как «нет данных».")
    r.italic, r.font.size = True, Pt(10)
    p.paragraph_format.line_spacing = 1.0


def iter_paragraphs(doc):
    def from_container(c):
        for p in c.paragraphs:
            yield p
        for t in getattr(c, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from from_container(cell)
    yield from from_container(doc)
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from from_container(part)


def build_from_template(template, cur, bases, cur_period, ctx, highlight, sources):
    doc = Document(template)
    table_slot = None
    for p in list(iter_paragraphs(doc)):
        if "{{" not in p.text:
            continue
        if p.text.strip() == "{{TABLE}}":
            table_slot = p
            continue
        like = p.runs[0] if p.runs else None
        segs = render(p.text, ctx)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        add_runs(p, segs, highlight, like)
    if table_slot is not None:
        t = summary_table(doc, cur, bases, cur_period)
        table_slot._p.addnext(t._tbl)
        table_slot._p.getparent().remove(table_slot._p)
    return doc

# =============================================================================
# 5. ЗАПУСК
# =============================================================================
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--current", required=True, help="итоговый отчёт за отчётный период")
    ap.add_argument("--previous", help="отчёт за предыдущий период")
    ap.add_argument("--appg", help="отчёт за аналогичный период прошлого года")
    ap.add_argument("--search", nargs="*", default=[], help="папки, где искать предыдущий период и АППГ")
    ap.add_argument("--template", help="образец доклада .docx с метками {{...}}")
    ap.add_argument("--out", help="куда сохранить .docx")
    ap.add_argument("--column", default=COLUMN_HINT, help="фрагмент названия графы")
    ap.add_argument("--period", help="отчётный период ГГГГ-ММ")
    ap.add_argument("--no-highlight", action="store_true", help="не подсвечивать подстановки")
    ap.add_argument("--list-labels", action="store_true", help="показать строки и значения из файла и выйти")
    a = ap.parse_args(argv)

    cur_rep = Report(a.current, a.column)
    if a.list_labels:
        print(f"Графа: {cur_rep.col_title}")
        for row in cur_rep.rows:
            print(f"{row['row_no']:>5}  {cur_rep.cell(row['r'], cur_rep.col)!s:>12}  {row['label']}")
        return 0
    if a.period:
        pm = re.match(r"^(\d{4})-(\d{1,2})$", a.period)
        if not pm:
            sys.exit("--period указывается как ГГГГ-ММ, например 2026-03")
        cur_rep.period = (int(pm.group(1)), int(pm.group(2)))
    if not cur_rep.period:
        sys.exit("Период не определён: укажите --period ГГГГ-ММ или назовите файл как 1175_230_ГГММ.XLS")
    y, m = cur_rep.period
    m_name = re.match(r"^\s*(\d+)_(\d+)_", os.path.basename(a.current))
    code, form = (m_name.group(1), m_name.group(2)) if m_name else ("1175", "230")

    prev_path, appg_path = a.previous, a.appg
    if a.search:
        if not prev_path:
            prev_path = next((p for mm in range(m - 1, 0, -1)
                              if (p := find_report(a.search, code, form, y, mm))), None)
        if not appg_path:
            appg_path = find_report(a.search, code, form, y - 1, m)

    cur = cur_rep.values()
    bases, sources, warnings = {}, [os.path.basename(a.current)], list(cur_rep.warnings)
    for tag, path, title in (("prev", prev_path, "предыдущий период"), ("appg", appg_path, "АППГ")):
        if not path:
            warnings.append(f"Нет отчёта за {title} — сравнение не выполняется.")
            continue
        rep = Report(path, a.column)
        bases[tag] = (rep.values(), rep.period)
        sources.append(os.path.basename(path))
        warnings += [f"[{title}] {w}" for w in rep.warnings]

    ctx = build_context(cur, (y, m), bases)
    highlight = not a.no_highlight
    if a.template:
        doc = build_from_template(a.template, cur, bases, (y, m), ctx, highlight, sources)
    else:
        doc = build_default(cur, bases, (y, m), ctx, highlight, sources)
    out = a.out or f"Доклад_{y}-{m:02d}.docx"
    doc.save(out)

    print(f"Период: {period_label(y, m)}; графа: {cur_rep.col_title}")
    for k, _, unit, title in INDICATORS:
        c = cur[k]
        extra = "  ".join(f"{t}: {fnum(bases[t][0][k]['value'])}" for t in bases)
        print(f"  {title[:55]:55} {fnum(c['value']):>12}  (стр. {c['row_no']})  {extra}")
    for w in warnings:
        print("ВНИМАНИЕ:", w)
    if UNKNOWN:
        print("ВНИМАНИЕ: в шаблоне неизвестные метки:", ", ".join(sorted(UNKNOWN)))
    print("Готово:", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
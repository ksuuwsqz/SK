"""
Экспорт черновика докладной в .docx через python-docx.
"""
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt


def export_report_to_docx(text: str, output_path: str, title: str = "Черновик докладной записки") -> str:
    """
    Сохраняет переданный текст в .docx.
    Возвращает фактический путь сохранения.
    """
    if not text.strip():
        raise ValueError("Текст докладной пуст — нечего экспортировать.")

    path = Path(output_path)
    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    if "ТЕСТ" in text.upper():
        doc.add_paragraph("ДОКУМЕНТ СОДЕРЖИТ ТЕСТОВЫЕ ДАННЫЕ.")

    for block in text.split("\n\n"):
        block = block.strip()
        if block:
            doc.add_paragraph(block)

    doc.save(path)
    return str(path)